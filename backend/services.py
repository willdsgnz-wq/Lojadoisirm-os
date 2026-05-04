from __future__ import annotations

import io
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook

from backend.auth import hash_password, verify_password
from backend.db import ForeignKeyViolation, UniqueViolation, get_connection
from backend.fiscal import DanfeService, XmlStorageService, get_fiscal_provider
from backend.product_importer import load_spreadsheet_rows
from backend.utils import ensure_date, iso_now, parse_iso_date, round_money, today_iso


PAYMENT_METHODS = ["Dinheiro", "Pix", "Débito", "Crédito", "Cheque", "Boleto", "À Prazo", "Outro", "Outros"]
SALES_PAYMENT_METHODS = ["À Prazo", "Boleto", "Dinheiro", "Pix", "Débito", "Crédito", "Cheque"]
PAYMENT_METHOD_ALIASES = {
    "dinheiro": "Dinheiro",
    "pix": "Pix",
    "debito": "Débito",
    "débito": "Débito",
    "credito": "Crédito",
    "crédito": "Crédito",
    "cheque": "Cheque",
    "boleto": "Boleto",
    "prazo": "À Prazo",
    "a prazo": "À Prazo",
    "à prazo": "À Prazo",
    "outro": "Outro",
    "outros": "Outros",
}
PRODUCT_UNITS = ["UN", "KG", "M3", "M2", "SC", "RL", "CH", "LT", "CX"]
QUOTE_ITEM_UNITS = ["UN", "MT", "M²", "M³", "KG", "SC", "CX", "PCT", "LT", "Outro"]
QUOTE_STATUSES = ["Pendente", "Aprovado", "Cancelado", "Nao aprovado"]
CHECK_STATUSES = ["Pendente", "Compensado", "Atrasado", "Cancelado"]
BILL_STATUSES = ["Pendente", "Pago", "Vencendo hoje", "Atrasado"]
STOCK_MOVEMENT_TYPES = ["ENTRADA", "SAIDA", "AJUSTE"]
FISCAL_ENVIRONMENTS = ["homologation", "production"]
FISCAL_PROVIDER_OPTIONS = ["mock", "focus_nfe", "nfe_io", "tecnospeed"]
CUSTOMER_PERSON_TYPES = ["PF", "PJ"]
CUSTOMER_IE_INDICATORS = ["Nao contribuinte", "Isento", "Contribuinte"]
DEFAULT_EXPENSE_CATEGORY = "Conta paga"
AUTO_BILL_EXPENSE_CATEGORY = "Boleto"
AUTO_BILL_EXPENSE_DESCRIPTION = "BOLETO"
AUTO_BILL_EXPENSE_PAYMENT_METHOD = "Boleto"

PRODUCT_IMPORT_MAPPING = {
    "sku": "sku",
    "categoria": "category",
    "produto": "name",
    "unidade": "unit",
    "preco venda": "sale_price",
    "estoque inicial": "stock_quantity",
    "ncm": "ncm",
    "cfop": "cfop_default",
    "origem": "origin",
    "csosn": "csosn",
    "ativo": "active",
    "observacoes": "notes",
}
PRODUCT_IMPORT_REQUIRED_COLUMNS = ["sku", "categoria", "produto", "unidade", "preco venda"]


class ServiceError(Exception):
    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.message = message
        self.status_code = status_code


def _clean_text(value: Any) -> str:
    return str(value or "").strip()


def _normalize_lookup_text(value: Any) -> str:
    cleaned = _clean_text(value)
    normalized = unicodedata.normalize("NFD", cleaned)
    ascii_only = "".join(char for char in normalized if unicodedata.category(char) != "Mn")
    return "".join(char for char in ascii_only.lower() if char.isalnum())


def _normalize_check_number(value: Any) -> str:
    cleaned = _require_text(value, "check_number")
    normalized = _normalize_lookup_text(cleaned)
    if normalized in {"sn", "semnumero"}:
        return "S/N"
    return cleaned


def _normalize_payment_method(value: Any, default: str = "") -> str:
    cleaned = _clean_text(value)
    if not cleaned:
        return default

    normalized_key = " ".join(cleaned.split()).casefold()
    return PAYMENT_METHOD_ALIASES.get(normalized_key, cleaned)


def _clean_digits(value: Any) -> str:
    return "".join(char for char in str(value or "") if char.isdigit())


def _require_text(value: Any, label: str) -> str:
    cleaned = _clean_text(value)
    if not cleaned:
        raise ServiceError(f"O campo '{label}' é obrigatório.")
    return cleaned


def _parse_amount(value: Any, label: str, *, min_value: float = 0, allow_zero: bool = True) -> float:
    try:
        numeric = float(value)
    except (TypeError, ValueError) as exc:
        raise ServiceError(f"Informe um valor válido para '{label}'.") from exc

    if numeric < min_value or (not allow_zero and numeric == 0):
        raise ServiceError(f"Informe um valor válido para '{label}'.")
    return round_money(numeric)


def _parse_int(value: Any, label: str, *, min_value: int | None = None, default: int | None = None) -> int:
    if value in (None, "") and default is not None:
        return default
    try:
        numeric = int(value)
    except (TypeError, ValueError) as exc:
        raise ServiceError(f"Informe um valor inteiro válido para '{label}'.") from exc
    if min_value is not None and numeric < min_value:
        raise ServiceError(f"Informe um valor válido para '{label}'.")
    return numeric


def _parse_bool(value: Any, default: bool = False) -> bool:
    if value in (None, ""):
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    normalized = str(value).strip().lower()
    if normalized in {"1", "true", "sim", "yes", "ativo", "s", "y"}:
        return True
    if normalized in {"0", "false", "nao", "não", "no", "inativo", "n"}:
        return False
    return default


def _normalize_product_unit(value: Any) -> str:
    normalized = _clean_text(value).upper()
    if normalized == "M²":
        normalized = "M2"
    if normalized == "M³":
        normalized = "M3"
    if not normalized:
        raise ServiceError("Selecione uma unidade válida para o produto.")
    if normalized not in PRODUCT_UNITS:
        raise ServiceError(f"Unidade inválida. Use uma destas opções: {', '.join(PRODUCT_UNITS)}.")
    return normalized


def _normalize_quote_unit(value: Any) -> str:
    normalized = _clean_text(value) or "UN"
    if normalized not in QUOTE_ITEM_UNITS:
        raise ServiceError(f"Selecione uma unidade válida no orçamento. Opções: {', '.join(QUOTE_ITEM_UNITS)}.")
    return normalized


def _normalize_sale_time(value: Any) -> str:
    time_value = _clean_text(value) or datetime.now().strftime("%H:%M")
    try:
        parsed = datetime.strptime(time_value, "%H:%M")
    except ValueError as exc:
        raise ServiceError("Informe um horário de venda válido no formato HH:MM.") from exc
    return parsed.strftime("%H:%M")


def _infer_period(time_value: str) -> str:
    return "Manhã" if time_value <= "12:00" else "Tarde"


def _rows_to_dicts(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [dict(row) for row in rows]


def _serialize_user(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": row["id"],
        "username": row["username"],
        "full_name": row["full_name"],
    }


def _serialize_product(row: dict[str, Any]) -> dict[str, Any]:
    item = dict(row)
    item["sku"] = item.get("sku") or item.get("code")
    item["code"] = item.get("code") or item["sku"]
    item["category"] = item.get("category") or ""
    item["name"] = item.get("name") or ""
    item["description"] = item.get("description") or ""
    item["notes"] = item.get("notes") or ""
    item["unit"] = _clean_text(item.get("unit")).upper() or "UN"
    item["active"] = bool(item.get("active", True))
    item["low_stock"] = float(item.get("stock_quantity") or 0) <= float(item.get("min_stock") or 0)
    item["out_of_stock"] = float(item.get("stock_quantity") or 0) <= 0

    # Aliases amigáveis para os módulos novos.
    item["nome_produto"] = item["name"]
    item["descricao_curta"] = item["description"]
    item["preco_venda"] = item["sale_price"]
    item["preco_custo"] = item["cost_price"]
    item["estoque_atual"] = item["stock_quantity"]
    item["estoque_minimo"] = item["min_stock"]
    item["cfop_padrao"] = item.get("cfop_default")
    item["observacoes"] = item["notes"]
    item["categoria"] = item["category"]
    return item


def _serialize_check(row: dict[str, Any]) -> dict[str, Any]:
    item = dict(row)
    status = item["status"]
    today = date.today()
    issue_date = parse_iso_date(item["issue_date"])
    due_date = parse_iso_date(item["due_date"])

    days_pending = (today - issue_date).days if issue_date else 0
    days_pending = max(days_pending, 0)

    days_overdue = 0
    effective_status = status
    if status not in {"Compensado", "Cancelado"} and due_date and today > due_date:
        effective_status = "Atrasado"
        days_overdue = (today - due_date).days

    item["effective_status"] = effective_status
    item["days_pending"] = days_pending
    item["days_overdue"] = days_overdue
    item["is_overdue"] = effective_status == "Atrasado"
    return item


def _serialize_bill(row: dict[str, Any]) -> dict[str, Any]:
    item = dict(row)
    due_date = parse_iso_date(item["due_date"])
    today = date.today()
    is_paid = bool(item.get("is_paid"))

    effective_status = "Pago" if is_paid else "Pendente"
    is_due_today = bool(not is_paid and due_date and due_date == today)
    is_overdue = bool(not is_paid and due_date and due_date < today)

    if is_overdue:
        effective_status = "Atrasado"
    elif is_due_today:
        effective_status = "Vencendo hoje"

    item["is_paid"] = is_paid
    item["effective_status"] = effective_status
    item["status"] = effective_status
    item["is_due_today"] = is_due_today
    item["is_overdue"] = is_overdue
    item["is_pending"] = not is_paid
    item["days_overdue"] = (today - due_date).days if is_overdue and due_date else 0
    return item


def _serialize_stock_movement(row: dict[str, Any]) -> dict[str, Any]:
    item = dict(row)
    item["movement_date"] = str(item.get("created_at") or "")[:10]
    item["product_name"] = item.get("product_name") or "Produto removido"
    item["product_sku"] = item.get("product_sku") or "-"
    return item


def _serialize_nfe_row(row: dict[str, Any]) -> dict[str, Any]:
    item = dict(row)
    item["source_type"] = _clean_text(item.get("source_type")) or ("sale" if item.get("sale_id") else "manual")
    item["customer_name"] = item.get("customer_name_resolved") or item.get("customer_name") or "Cliente não informado"
    item["payment_method"] = _normalize_payment_method(
        item.get("payment_method_resolved") or item.get("payment_method"),
        default="-",
    ) or "-"
    item["sale_total_amount"] = round_money(item.get("sale_total_amount") or 0)
    item["total_amount"] = round_money(item.get("total_amount_resolved") or item.get("total_amount") or item["sale_total_amount"])
    item["has_xml"] = bool(item.get("xml_path"))
    item["has_pdf"] = bool(item.get("pdf_path"))
    return item


def _serialize_expense_row(row: dict[str, Any]) -> dict[str, Any]:
    item = dict(row)
    item["payment_method"] = _normalize_payment_method(item.get("payment_method"), default="-") or "-"
    item.pop("category", None)
    item.pop("supplier", None)
    item.pop("notes", None)
    return item


def _payment_date_from_paid_at(value: Any) -> str:
    cleaned = _clean_text(value)
    return cleaned[:10] if cleaned else today_iso()


def _get_linked_bill_expense(connection: Any, bill_id: int) -> dict[str, Any] | None:
    row = connection.execute("SELECT * FROM expenses WHERE linked_bill_id = ?", (bill_id,)).fetchone()
    return dict(row) if row else None


def _sync_paid_bill_expense(
    connection: Any,
    *,
    bill_id: int,
    is_paid: bool,
    amount: float,
    payment_date: str,
) -> None:
    linked_expense = _get_linked_bill_expense(connection, bill_id)

    if not is_paid:
        if linked_expense:
            connection.execute("DELETE FROM expenses WHERE id = ?", (linked_expense["id"],))
        return

    payload = (
        payment_date,
        AUTO_BILL_EXPENSE_DESCRIPTION,
        AUTO_BILL_EXPENSE_CATEGORY,
        amount,
        AUTO_BILL_EXPENSE_PAYMENT_METHOD,
        "",
        "",
    )

    if linked_expense:
        connection.execute(
            """
            UPDATE expenses
            SET payment_date = ?, description = ?, category = ?, amount = ?,
                payment_method = ?, supplier = ?, notes = ?
            WHERE id = ?
            """,
            (*payload, linked_expense["id"]),
        )
        return

    connection.execute(
        """
        INSERT INTO expenses (
            payment_date, description, category, amount,
            payment_method, supplier, notes, linked_bill_id
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (*payload, bill_id),
    )


def _load_products_map(connection: Any, product_ids: list[int]) -> dict[int, dict[str, Any]]:
    if not product_ids:
        return {}
    placeholders = ",".join("?" for _ in product_ids)
    rows = connection.execute(
        f"SELECT * FROM products WHERE id IN ({placeholders})",
        tuple(product_ids),
    ).fetchall()
    return {row["id"]: dict(row) for row in rows}


def _current_sale_date() -> str:
    return datetime.now().strftime("%Y-%m-%d")


def _default_customer_name(row: dict[str, Any]) -> str:
    return row.get("customer_name") or "Consumidor final"


def _format_customer_document(value: Any) -> str:
    digits = _clean_digits(value)
    if len(digits) == 11:
        return f"{digits[:3]}.{digits[3:6]}.{digits[6:9]}-{digits[9:]}"
    if len(digits) == 14:
        return f"{digits[:2]}.{digits[2:5]}.{digits[5:8]}/{digits[8:12]}-{digits[12:]}"
    return _clean_text(value)


def _format_zip_code(value: Any) -> str:
    digits = _clean_digits(value)
    if len(digits) == 8:
        return f"{digits[:5]}-{digits[5:]}"
    return _clean_text(value)


def _normalize_customer_person_type(value: Any, document_digits: str = "", *, strict: bool = False) -> str:
    normalized = _normalize_lookup_text(value)
    mapping = {
        "pf": "PF",
        "pessoafisica": "PF",
        "fisica": "PF",
        "pj": "PJ",
        "pessoajuridica": "PJ",
        "juridica": "PJ",
    }

    if not normalized:
        if len(document_digits) == 14:
            return "PJ"
        return "PF"

    person_type = mapping.get(normalized)
    if person_type:
        return person_type
    if strict:
        raise ServiceError("Selecione um tipo de cliente valido: Pessoa Fisica ou Pessoa Juridica.")
    return "PF"


def _normalize_customer_ie_indicator(value: Any, *, strict: bool = False) -> str:
    normalized = _normalize_lookup_text(value)
    mapping = {
        "naocontribuinte": "Nao contribuinte",
        "naocontribuinteicms": "Nao contribuinte",
        "isento": "Isento",
        "contribuinte": "Contribuinte",
    }
    if not normalized:
        return "Nao contribuinte"
    indicator = mapping.get(normalized)
    if indicator:
        return indicator
    if strict:
        raise ServiceError("Selecione um indicador de IE valido.")
    return "Nao contribuinte"


def _compose_customer_address(customer: dict[str, Any]) -> str:
    street = _clean_text(customer.get("street"))
    number = _clean_text(customer.get("number"))
    complement = _clean_text(customer.get("complement"))
    district = _clean_text(customer.get("district"))
    city = _clean_text(customer.get("city"))
    state = _clean_text(customer.get("state")).upper()
    zip_code = _format_zip_code(customer.get("zip_code"))

    line_one = ", ".join(part for part in [street, number] if part)
    if complement:
        line_one = ", ".join(part for part in [line_one, complement] if part)
    city_state = "/".join(part for part in [city, state] if part)

    parts = [line_one, district, city_state]
    if zip_code:
        parts.append(f"CEP {zip_code}")
    return ", ".join(part for part in parts if part)


def _serialize_customer(row: dict[str, Any]) -> dict[str, Any]:
    item = dict(row)
    legacy_digits = _clean_digits(item.get("document"))
    cpf = _clean_digits(item.get("cpf"))
    cnpj = _clean_digits(item.get("cnpj"))
    if not cpf and not cnpj:
        if len(legacy_digits) == 11:
            cpf = legacy_digits
        elif len(legacy_digits) == 14:
            cnpj = legacy_digits

    item["person_type"] = _normalize_customer_person_type(item.get("person_type"), cpf or cnpj or legacy_digits)
    item["cpf"] = cpf
    item["cnpj"] = cnpj
    item["document"] = cpf or cnpj or legacy_digits
    item["document_formatted"] = _format_customer_document(item["document"])
    item["trade_name"] = _clean_text(item.get("trade_name"))
    item["phone"] = _clean_text(item.get("phone"))
    item["whatsapp"] = _clean_text(item.get("whatsapp"))
    item["email"] = _clean_text(item.get("email"))
    item["zip_code"] = _clean_digits(item.get("zip_code"))
    item["zip_code_formatted"] = _format_zip_code(item["zip_code"])
    item["street"] = _clean_text(item.get("street"))
    item["number"] = _clean_text(item.get("number"))
    item["complement"] = _clean_text(item.get("complement"))
    item["district"] = _clean_text(item.get("district"))
    item["city"] = _clean_text(item.get("city"))
    item["state"] = _clean_text(item.get("state")).upper()
    item["city_ibge_code"] = _clean_digits(item.get("city_ibge_code"))
    item["ie_indicator"] = _normalize_customer_ie_indicator(item.get("ie_indicator"))
    item["state_registration"] = _clean_text(item.get("state_registration"))
    item["rg"] = _clean_text(item.get("rg"))
    item["birth_date"] = _clean_text(item.get("birth_date"))
    item["notes"] = _clean_text(item.get("notes"))
    if item["person_type"] != "PF":
        item["rg"] = ""
        item["birth_date"] = ""
    if item["ie_indicator"] != "Contribuinte":
        item["state_registration"] = ""
    item["address"] = _compose_customer_address(item) or _clean_text(item.get("address"))
    item["name"] = _clean_text(item.get("name"))
    item["is_company"] = item["person_type"] == "PJ"
    item["requires_state_registration"] = item["ie_indicator"] == "Contribuinte"
    item["city_label"] = " / ".join(part for part in [item["city"], item["state"]] if part)
    return item


def _validate_cpf(value: Any) -> str:
    digits = _clean_digits(value)
    if len(digits) != 11 or digits == digits[0] * 11:
        raise ServiceError("Informe um CPF valido com 11 digitos.")

    for digit_index in range(9, 11):
        total = sum(int(digits[position]) * ((digit_index + 1) - position) for position in range(digit_index))
        expected = (total * 10) % 11
        expected = 0 if expected == 10 else expected
        if expected != int(digits[digit_index]):
            raise ServiceError("Informe um CPF valido.")
    return digits


def _ensure_customer_document_not_duplicated(
    connection: Any,
    *,
    cpf: str = "",
    cnpj: str = "",
    exclude_id: int | None = None,
) -> None:
    if cpf:
        params: list[Any] = [cpf]
        query = "SELECT id, name FROM customers WHERE cpf = ?"
        if exclude_id:
            query += " AND id <> ?"
            params.append(exclude_id)
        row = connection.execute(query + " LIMIT 1", tuple(params)).fetchone()
        if row:
            raise ServiceError(f"Ja existe um cliente cadastrado com esse CPF: {row['name']}.", 409)

    if cnpj:
        params = [cnpj]
        query = "SELECT id, name FROM customers WHERE cnpj = ?"
        if exclude_id:
            query += " AND id <> ?"
            params.append(exclude_id)
        row = connection.execute(query + " LIMIT 1", tuple(params)).fetchone()
        if row:
            raise ServiceError(f"Ja existe um cliente cadastrado com esse CNPJ: {row['name']}.", 409)


def _build_customer_values(payload: dict[str, Any], current: dict[str, Any] | None = None) -> dict[str, Any]:
    current = current or {}
    legacy_document = payload.get("document") if payload.get("document") is not None else current.get("document")
    legacy_digits = _clean_digits(legacy_document)
    raw_cpf = payload.get("cpf") if payload.get("cpf") is not None else current.get("cpf")
    raw_cnpj = payload.get("cnpj") if payload.get("cnpj") is not None else current.get("cnpj")
    cpf_digits = _clean_digits(raw_cpf)
    cnpj_digits = _clean_digits(raw_cnpj)

    if not cpf_digits and not cnpj_digits and legacy_digits:
        if len(legacy_digits) == 11:
            cpf_digits = legacy_digits
        elif len(legacy_digits) == 14:
            cnpj_digits = legacy_digits

    person_type = _normalize_customer_person_type(
        payload.get("person_type") if payload.get("person_type") is not None else current.get("person_type"),
        cpf_digits or cnpj_digits or legacy_digits,
        strict=True,
    )

    name_label = "nome_completo" if person_type == "PF" else "razao_social"
    name = _require_text(payload.get("name") if payload.get("name") is not None else current.get("name"), name_label)
    trade_name = _clean_text(payload.get("trade_name") if payload.get("trade_name") is not None else current.get("trade_name"))
    phone = _clean_text(payload.get("phone") if payload.get("phone") is not None else current.get("phone"))
    whatsapp = _clean_text(payload.get("whatsapp") if payload.get("whatsapp") is not None else current.get("whatsapp"))
    email = _clean_text(payload.get("email") if payload.get("email") is not None else current.get("email"))
    if email and ("@" not in email or "." not in email.split("@")[-1]):
        raise ServiceError("Informe um e-mail valido para o cliente.")

    zip_code = _clean_digits(payload.get("zip_code") if payload.get("zip_code") is not None else current.get("zip_code"))
    if len(zip_code) != 8:
        raise ServiceError("Informe um CEP valido com 8 digitos.")

    street = _require_text(payload.get("street") if payload.get("street") is not None else current.get("street"), "logradouro")
    number = _require_text(payload.get("number") if payload.get("number") is not None else current.get("number"), "numero")
    complement = _clean_text(payload.get("complement") if payload.get("complement") is not None else current.get("complement"))
    district = _require_text(payload.get("district") if payload.get("district") is not None else current.get("district"), "bairro")
    city = _require_text(payload.get("city") if payload.get("city") is not None else current.get("city"), "cidade")
    state = _clean_text(payload.get("state") if payload.get("state") is not None else current.get("state")).upper()
    if len(state) != 2 or not state.isalpha():
        raise ServiceError("Informe uma UF valida com 2 letras.")

    city_ibge_code = _clean_digits(
        payload.get("city_ibge_code") if payload.get("city_ibge_code") is not None else current.get("city_ibge_code")
    )
    if len(city_ibge_code) != 7:
        raise ServiceError("Informe o codigo IBGE do municipio com 7 digitos.")

    ie_indicator = _normalize_customer_ie_indicator(
        payload.get("ie_indicator") if payload.get("ie_indicator") is not None else current.get("ie_indicator"),
        strict=True,
    )
    state_registration = _clean_text(
        payload.get("state_registration")
        if payload.get("state_registration") is not None
        else current.get("state_registration")
    )
    if ie_indicator == "Contribuinte" and not state_registration:
        raise ServiceError("A inscricao estadual e obrigatoria para cliente contribuinte.")
    if ie_indicator != "Contribuinte":
        state_registration = ""

    if person_type == "PF":
        cpf_digits = _validate_cpf(cpf_digits or legacy_digits)
        cnpj_digits = ""
        rg = _clean_text(payload.get("rg") if payload.get("rg") is not None else current.get("rg"))
        birth_date = _clean_text(payload.get("birth_date") if payload.get("birth_date") is not None else current.get("birth_date"))
        if birth_date:
            birth_date = ensure_date(birth_date)
        trade_name = ""
    else:
        cnpj_digits = _validate_cnpj(cnpj_digits or legacy_digits)
        cpf_digits = ""
        rg = ""
        birth_date = ""

    values = {
        "person_type": person_type,
        "name": name,
        "trade_name": trade_name,
        "cpf": cpf_digits,
        "cnpj": cnpj_digits,
        "phone": phone,
        "whatsapp": whatsapp,
        "email": email,
        "zip_code": zip_code,
        "street": street,
        "number": number,
        "complement": complement,
        "district": district,
        "city": city,
        "state": state,
        "city_ibge_code": city_ibge_code,
        "ie_indicator": ie_indicator,
        "state_registration": state_registration,
        "rg": rg,
        "birth_date": birth_date,
        "notes": _clean_text(payload.get("notes") if payload.get("notes") is not None else current.get("notes")),
    }
    values["document"] = cpf_digits or cnpj_digits
    values["address"] = _compose_customer_address(values)
    return values


def _get_customer_missing_fiscal_fields(customer: dict[str, Any]) -> list[str]:
    labels = [
        ("name", "Nome completo" if customer.get("person_type") == "PF" else "Razao social"),
        ("document", "CPF" if customer.get("person_type") == "PF" else "CNPJ"),
        ("zip_code", "CEP"),
        ("street", "Logradouro"),
        ("number", "Numero"),
        ("district", "Bairro"),
        ("city", "Cidade"),
        ("state", "UF"),
        ("city_ibge_code", "Codigo IBGE"),
        ("ie_indicator", "Indicador de IE"),
    ]
    missing = [label for field, label in labels if not _clean_text(customer.get(field))]
    if customer.get("ie_indicator") == "Contribuinte" and not _clean_text(customer.get("state_registration")):
        missing.append("Inscricao estadual")
    return missing


def _load_customer_for_nfe(connection: Any, customer_id: int) -> dict[str, Any]:
    row = connection.execute("SELECT * FROM customers WHERE id = ?", (customer_id,)).fetchone()
    if not row:
        raise ServiceError("Cliente nao encontrado.", 404)
    customer = _serialize_customer(row)
    missing = _get_customer_missing_fiscal_fields(customer)
    if missing:
        raise ServiceError(
            "Complete o cadastro do cliente antes de emitir a NF-e: " + ", ".join(missing) + ".",
            409,
        )
    return customer


def ensure_demo_user() -> None:
    with get_connection() as connection:
        existing = connection.execute("SELECT id FROM users WHERE username = ?", ("admin",)).fetchone()
        if existing:
            return
        connection.execute(
            """
            INSERT INTO users (username, full_name, password_hash)
            VALUES (?, ?, ?)
            """,
            ("admin", "Administrador da Loja", hash_password("123456")),
        )


def authenticate_user(username: str, password: str) -> dict[str, Any]:
    with get_connection() as connection:
        row = connection.execute("SELECT * FROM users WHERE username = ?", (username.strip(),)).fetchone()
    if not row or not verify_password(password, row["password_hash"]):
        raise ServiceError("Usuário ou senha inválidos.", 401)
    return _serialize_user(row)


def get_user_by_id(user_id: int | str | None) -> dict[str, Any] | None:
    if not user_id:
        return None
    with get_connection() as connection:
        row = connection.execute("SELECT * FROM users WHERE id = ?", (int(user_id),)).fetchone()
    if not row:
        return None
    return _serialize_user(row)


def list_products() -> list[dict[str, Any]]:
    with get_connection() as connection:
        rows = connection.execute(
            """
            SELECT *
            FROM products
            ORDER BY active DESC, LOWER(name), id DESC
            """
        ).fetchall()
    return [_serialize_product(row) for row in rows]


def get_product(product_id: int) -> dict[str, Any]:
    with get_connection() as connection:
        row = connection.execute("SELECT * FROM products WHERE id = ?", (product_id,)).fetchone()
    if not row:
        raise ServiceError("Produto não encontrado.", 404)
    return _serialize_product(row)


def _insert_stock_movement(
    connection: Any,
    *,
    product_id: int,
    movement_type: str,
    quantity: float,
    balance_before: float,
    balance_after: float,
    reason: str,
    document_reference: str = "",
    user_id: int | None = None,
) -> None:
    connection.execute(
        """
        INSERT INTO stock_movements (
            product_id, movement_type, quantity, balance_before, balance_after,
            reason, document_reference, user_id, created_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            product_id,
            movement_type,
            round_money(quantity),
            round_money(balance_before),
            round_money(balance_after),
            _clean_text(reason),
            _clean_text(document_reference),
            user_id,
            iso_now(),
        ),
    )


def _allow_negative_stock(connection: Any) -> bool:
    row = connection.execute(
        """
        SELECT allow_negative_stock
        FROM fiscal_settings
        ORDER BY id ASC
        LIMIT 1
        """
    ).fetchone()
    return bool(row and row.get("allow_negative_stock"))


def create_product(payload: dict[str, Any]) -> dict[str, Any]:
    sku = _require_text(payload.get("sku") or payload.get("code"), "sku")
    name = _require_text(payload.get("name") or payload.get("nome_produto"), "name")
    category = _require_text(payload.get("category") or payload.get("categoria"), "category")
    unit = _normalize_product_unit(payload.get("unit") or payload.get("unidade"))
    cost_price = _parse_amount(payload.get("cost_price", payload.get("preco_custo", 0)), "cost_price", min_value=0)
    sale_price = _parse_amount(payload.get("sale_price", payload.get("preco_venda", 0)), "sale_price", min_value=0)
    stock_quantity = _parse_amount(payload.get("stock_quantity", payload.get("estoque_atual", 0)), "stock_quantity", min_value=0)
    min_stock = _parse_amount(payload.get("min_stock", payload.get("estoque_minimo", 0)), "min_stock", min_value=0)
    now = iso_now()

    try:
        with get_connection() as connection:
            cursor = connection.execute(
                """
                INSERT INTO products (
                    sku, code, category, name, description, unit, cost_price, sale_price,
                    stock_quantity, min_stock, ncm, cfop_default, origin, csosn, notes, active,
                    updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    sku,
                    _clean_text(payload.get("code")) or sku,
                    category,
                    name,
                    _clean_text(payload.get("description") or payload.get("descricao_curta")),
                    unit,
                    cost_price,
                    sale_price,
                    stock_quantity,
                    min_stock,
                    _clean_text(payload.get("ncm")),
                    _clean_text(payload.get("cfop_default") or payload.get("cfop_padrao")),
                    _clean_text(payload.get("origin") or payload.get("origem")),
                    _clean_text(payload.get("csosn")),
                    _clean_text(payload.get("notes") or payload.get("observacoes")),
                    _parse_bool(payload.get("active") if payload.get("active") is not None else payload.get("ativo"), True),
                    now,
                ),
            )
            product_id = cursor.lastrowid

            if stock_quantity > 0:
                _insert_stock_movement(
                    connection,
                    product_id=product_id,
                    movement_type="ENTRADA",
                    quantity=stock_quantity,
                    balance_before=0,
                    balance_after=stock_quantity,
                    reason="Estoque inicial do cadastro do produto",
                    document_reference=f"PRODUTO-{product_id}",
                    user_id=_parse_int(payload.get("_user_id"), "user_id", default=0) if payload.get("_user_id") else None,
                )
    except UniqueViolation as exc:
        raise ServiceError("Já existe um produto com esse SKU/código.") from exc

    return get_product(product_id)


def update_product(product_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    current = get_product(product_id)
    values = {
        "sku": _require_text(payload.get("sku", current["sku"]) or payload.get("code", current["code"]), "sku"),
        "code": _clean_text(payload.get("code", current["code"])) or current["sku"],
        "category": _require_text(payload.get("category", current["category"]), "category"),
        "name": _require_text(payload.get("name", current["name"]), "name"),
        "description": _clean_text(payload.get("description") if payload.get("description") is not None else current["description"]),
        "unit": _normalize_product_unit(payload.get("unit", current["unit"])),
        "cost_price": _parse_amount(payload.get("cost_price", current["cost_price"]), "cost_price", min_value=0),
        "sale_price": _parse_amount(payload.get("sale_price", current["sale_price"]), "sale_price", min_value=0),
        "stock_quantity": _parse_amount(payload.get("stock_quantity", current["stock_quantity"]), "stock_quantity", min_value=0),
        "min_stock": _parse_amount(payload.get("min_stock", current["min_stock"]), "min_stock", min_value=0),
        "ncm": _clean_text(payload.get("ncm") if payload.get("ncm") is not None else current.get("ncm")),
        "cfop_default": _clean_text(payload.get("cfop_default") if payload.get("cfop_default") is not None else current.get("cfop_default")),
        "origin": _clean_text(payload.get("origin") if payload.get("origin") is not None else current.get("origin")),
        "csosn": _clean_text(payload.get("csosn") if payload.get("csosn") is not None else current.get("csosn")),
        "notes": _clean_text(payload.get("notes") if payload.get("notes") is not None else current.get("notes")),
        "active": _parse_bool(payload.get("active"), current["active"]),
        "updated_at": iso_now(),
    }

    try:
        with get_connection() as connection:
            before_stock = float(current["stock_quantity"])
            after_stock = float(values["stock_quantity"])

            connection.execute(
                """
                UPDATE products
                SET sku = ?, code = ?, category = ?, name = ?, description = ?, unit = ?, cost_price = ?,
                    sale_price = ?, stock_quantity = ?, min_stock = ?, ncm = ?, cfop_default = ?,
                    origin = ?, csosn = ?, notes = ?, active = ?, updated_at = ?
                WHERE id = ?
                """,
                (
                    values["sku"],
                    values["code"],
                    values["category"],
                    values["name"],
                    values["description"],
                    values["unit"],
                    values["cost_price"],
                    values["sale_price"],
                    values["stock_quantity"],
                    values["min_stock"],
                    values["ncm"],
                    values["cfop_default"],
                    values["origin"],
                    values["csosn"],
                    values["notes"],
                    values["active"],
                    values["updated_at"],
                    product_id,
                ),
            )

            if round_money(after_stock - before_stock) != 0:
                _insert_stock_movement(
                    connection,
                    product_id=product_id,
                    movement_type="AJUSTE",
                    quantity=after_stock - before_stock,
                    balance_before=before_stock,
                    balance_after=after_stock,
                    reason="Ajuste realizado pelo cadastro do produto",
                    document_reference=f"PRODUTO-{product_id}",
                    user_id=_parse_int(payload.get("_user_id"), "user_id", default=0) if payload.get("_user_id") else None,
                )
    except UniqueViolation as exc:
        raise ServiceError("Já existe outro produto com esse SKU/código.") from exc

    return get_product(product_id)


def delete_product(product_id: int) -> None:
    try:
        with get_connection() as connection:
            deleted = connection.execute("DELETE FROM products WHERE id = ?", (product_id,)).rowcount
    except ForeignKeyViolation as exc:
        raise ServiceError("Esse produto já foi usado em vendas. Desative o produto em vez de excluir.", 409) from exc
    if not deleted:
        raise ServiceError("Produto não encontrado.", 404)


def import_products_from_spreadsheet(filename: str, content: bytes) -> dict[str, Any]:
    rows = load_spreadsheet_rows(filename, content)
    if not rows:
        raise ServiceError("A planilha enviada não possui linhas de produtos para importar.")

    available_columns = set(rows[0].keys())
    missing_columns = [column for column in PRODUCT_IMPORT_REQUIRED_COLUMNS if column not in available_columns]
    if missing_columns:
        labels = ", ".join(missing_columns)
        raise ServiceError(f"A planilha não contém as colunas obrigatórias: {labels}.")

    seen_skus: set[str] = set()
    report = {
        "imported": 0,
        "updated": 0,
        "ignored": 0,
        "errors": [],
    }

    for index, row in enumerate(rows, start=2):
        sku = _clean_text(row.get("sku"))
        if not sku:
            report["errors"].append(f"Linha {index}: SKU não informado.")
            continue

        if sku in seen_skus:
            report["ignored"] += 1
            report["errors"].append(f"Linha {index}: SKU {sku} repetido na mesma planilha.")
            continue
        seen_skus.add(sku)

        mapped_payload: dict[str, Any] = {}
        for source_column, target_field in PRODUCT_IMPORT_MAPPING.items():
            mapped_payload[target_field] = row.get(source_column)

        mapped_payload["code"] = sku
        mapped_payload["sku"] = sku
        mapped_payload["cost_price"] = row.get("preco custo") or row.get("custo") or 0
        mapped_payload["min_stock"] = row.get("estoque minimo") or row.get("estoque mínimo") or 0
        mapped_payload["description"] = row.get("descricao curta") or row.get("descrição curta") or row.get("descricao") or ""

        try:
            existing = next((product for product in list_products() if product["sku"] == sku), None)
            if existing:
                update_product(existing["id"], mapped_payload)
                report["updated"] += 1
            else:
                create_product(mapped_payload)
                report["imported"] += 1
        except ServiceError as exc:
            report["errors"].append(f"Linha {index}: {exc.message}")

    return report


def export_products_dataset(export_format: str = "csv") -> dict[str, Any]:
    products = list_products()
    columns = [
        ("SKU", "sku"),
        ("Categoria", "category"),
        ("Produto", "name"),
        ("Descrição curta", "description"),
        ("Unidade", "unit"),
        ("Preço Venda", "sale_price"),
        ("Preço Custo", "cost_price"),
        ("Estoque Atual", "stock_quantity"),
        ("Estoque Mínimo", "min_stock"),
        ("NCM", "ncm"),
        ("CFOP", "cfop_default"),
        ("Origem", "origin"),
        ("CSOSN", "csosn"),
        ("Ativo", "active"),
        ("Observações", "notes"),
    ]

    if export_format == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Produtos"
        sheet.append([label for label, _key in columns])
        for product in products:
            sheet.append([product.get(key, "") for _label, key in columns])

        content = io.BytesIO()
        workbook.save(content)
        return {
            "filename": f"produtos-{today_iso()}.xlsx",
            "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "content": content.getvalue(),
        }

    lines = [";".join(label for label, _key in columns)]
    for product in products:
        lines.append(";".join(str(product.get(key, "")) for _label, key in columns))
    return {
        "filename": f"produtos-{today_iso()}.csv",
        "content_type": "text/csv; charset=utf-8",
        "content": ("\ufeff" + "\n".join(lines)).encode("utf-8"),
    }


def list_customers() -> list[dict[str, Any]]:
    with get_connection() as connection:
        rows = connection.execute(
            """
            SELECT *
            FROM customers
            ORDER BY LOWER(name), id DESC
            """
        ).fetchall()
    return [_serialize_customer(row) for row in rows]


def create_customer(payload: dict[str, Any]) -> dict[str, Any]:
    with get_connection() as connection:
        values = _build_customer_values(payload)
        _ensure_customer_document_not_duplicated(connection, cpf=values["cpf"], cnpj=values["cnpj"])
        try:
            cursor = connection.execute(
                """
                INSERT INTO customers (
                    person_type, name, trade_name, cpf, cnpj, phone, whatsapp, email, zip_code,
                    street, number, complement, district, city, state, city_ibge_code,
                    ie_indicator, state_registration, rg, birth_date, document, address, notes, updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    values["person_type"],
                    values["name"],
                    values["trade_name"] or None,
                    values["cpf"] or None,
                    values["cnpj"] or None,
                    values["phone"] or None,
                    values["whatsapp"] or None,
                    values["email"] or None,
                    values["zip_code"],
                    values["street"],
                    values["number"],
                    values["complement"] or None,
                    values["district"],
                    values["city"],
                    values["state"],
                    values["city_ibge_code"],
                    values["ie_indicator"],
                    values["state_registration"] or None,
                    values["rg"] or None,
                    values["birth_date"] or None,
                    values["document"],
                    values["address"],
                    values["notes"] or None,
                    iso_now(),
                ),
            )
            customer_id = cursor.lastrowid
        except UniqueViolation as exc:
            raise ServiceError("Ja existe um cliente com esse CPF/CNPJ.", 409) from exc
    return get_customer(customer_id)


def get_customer(customer_id: int) -> dict[str, Any]:
    with get_connection() as connection:
        row = connection.execute("SELECT * FROM customers WHERE id = ?", (customer_id,)).fetchone()
    if not row:
        raise ServiceError("Cliente nao encontrado.", 404)
    return _serialize_customer(row)


def update_customer(customer_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    current = get_customer(customer_id)
    with get_connection() as connection:
        values = _build_customer_values(payload, current)
        _ensure_customer_document_not_duplicated(
            connection,
            cpf=values["cpf"],
            cnpj=values["cnpj"],
            exclude_id=customer_id,
        )
        try:
            connection.execute(
                """
                UPDATE customers
                SET person_type = ?, name = ?, trade_name = ?, cpf = ?, cnpj = ?, phone = ?, whatsapp = ?, email = ?,
                    zip_code = ?, street = ?, number = ?, complement = ?, district = ?, city = ?, state = ?,
                    city_ibge_code = ?, ie_indicator = ?, state_registration = ?, rg = ?, birth_date = ?, document = ?, address = ?,
                    notes = ?, updated_at = ?
                WHERE id = ?
                """,
                (
                    values["person_type"],
                    values["name"],
                    values["trade_name"] or None,
                    values["cpf"] or None,
                    values["cnpj"] or None,
                    values["phone"] or None,
                    values["whatsapp"] or None,
                    values["email"] or None,
                    values["zip_code"],
                    values["street"],
                    values["number"],
                    values["complement"] or None,
                    values["district"],
                    values["city"],
                    values["state"],
                    values["city_ibge_code"],
                    values["ie_indicator"],
                    values["state_registration"] or None,
                    values["rg"] or None,
                    values["birth_date"] or None,
                    values["document"],
                    values["address"],
                    values["notes"] or None,
                    iso_now(),
                    customer_id,
                ),
            )
        except UniqueViolation as exc:
            raise ServiceError("Ja existe um cliente com esse CPF/CNPJ.", 409) from exc
    return get_customer(customer_id)


def delete_customer(customer_id: int) -> None:
    try:
        with get_connection() as connection:
            deleted = connection.execute("DELETE FROM customers WHERE id = ?", (customer_id,)).rowcount
    except ForeignKeyViolation as exc:
        raise ServiceError("Esse cliente ja foi usado em vendas ou orcamentos e nao pode ser excluido.", 409) from exc
    if not deleted:
        raise ServiceError("Cliente nao encontrado.", 404)


def _prepare_quote_items(raw_items: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], float]:
    if not raw_items:
        raise ServiceError("Adicione pelo menos um item ao orçamento.")

    items: list[dict[str, Any]] = []
    for index, item in enumerate(raw_items, start=1):
        product_id = (
            _parse_int(item.get("product_id"), f"product_id_{index}", min_value=1)
            if item.get("product_id") not in (None, "")
            else None
        )
        item_name = _clean_text(item.get("item_name"))
        if not item_name:
            raise ServiceError(f"Informe o nome do item {index}.")
        unit = _normalize_quote_unit(item.get("unit"))
        quantity = _parse_amount(item.get("quantity"), f"quantity_{index}", min_value=0.01, allow_zero=False)
        unit_price = _parse_amount(item.get("unit_price"), f"unit_price_{index}", min_value=0, allow_zero=True)
        total_price = round_money(quantity * unit_price)
        items.append(
            {
                "product_id": product_id,
                "item_name": item_name,
                "unit": unit,
                "quantity": quantity,
                "unit_price": unit_price,
                "total_price": total_price,
            }
        )

    subtotal_amount = round_money(sum(item["total_price"] for item in items))
    return items, subtotal_amount


def _serialize_quote_rows(rows: list[dict[str, Any]], items_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    items_by_quote: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for row in items_rows:
        items_by_quote[row["quote_id"]].append(
            {
                "id": row["id"],
                "product_id": row["product_id"],
                "item_name": row["item_name"],
                "product_name": row["item_name"],
                "unit": row["unit"],
                "quantity": row["quantity"],
                "unit_price": row["unit_price"],
                "total_price": row["total_price"],
            }
        )

    quotes: list[dict[str, Any]] = []
    for row in rows:
        quote = dict(row)
        quote["customer_name_manual"] = _clean_text(quote.get("customer_name_manual"))
        quote["customer_name"] = quote["customer_name"] or "Cliente não informado"
        quote["validity_date"] = quote.get("validity_date") or quote["quote_date"]
        quote["subtotal_amount"] = round_money(quote.get("subtotal_amount", quote["total_amount"]))
        quote["discount_amount"] = round_money(quote.get("discount_amount", 0))
        quote["items"] = items_by_quote.get(row["id"], [])
        quotes.append(quote)
    return quotes


def list_quotes() -> list[dict[str, Any]]:
    with get_connection() as connection:
        rows = connection.execute(
            """
            SELECT
                quotes.*,
                COALESCE(NULLIF(TRIM(quotes.customer_name_manual), ''), customers.name) AS customer_name
            FROM quotes
            LEFT JOIN customers ON customers.id = quotes.customer_id
            ORDER BY quotes.quote_date DESC, quotes.id DESC
            """
        ).fetchall()
        quote_ids = [row["id"] for row in rows]
        if quote_ids:
            placeholders = ",".join("?" for _ in quote_ids)
            items_rows = connection.execute(
                f"""
                SELECT quote_items.*
                FROM quote_items
                WHERE quote_items.quote_id IN ({placeholders})
                ORDER BY quote_items.id
                """,
                tuple(quote_ids),
            ).fetchall()
        else:
            items_rows = []
    return _serialize_quote_rows(rows, items_rows)


def create_quote(payload: dict[str, Any]) -> dict[str, Any]:
    quote_date = ensure_date(payload.get("quote_date"))
    validity_date = ensure_date(payload.get("validity_date"), quote_date)
    status = payload.get("status") or "Pendente"
    if status not in QUOTE_STATUSES:
        raise ServiceError("Escolha um status de orçamento válido.")

    customer_id = payload.get("customer_id") or None
    customer_name_manual = _clean_text(payload.get("customer_name_manual"))
    items, subtotal_amount = _prepare_quote_items(payload.get("items") or [])
    discount_amount = _parse_amount(payload.get("discount_amount", 0), "discount_amount", min_value=0, allow_zero=True)
    if discount_amount > subtotal_amount:
        raise ServiceError("O desconto não pode ser maior que o subtotal do orçamento.")
    total_amount = round_money(subtotal_amount - discount_amount)

    with get_connection() as connection:
        cursor = connection.execute(
            """
            INSERT INTO quotes (
                quote_date, validity_date, customer_id, customer_name_manual, subtotal_amount,
                discount_amount, total_amount, status, notes
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                quote_date,
                validity_date,
                int(customer_id) if customer_id else None,
                customer_name_manual or None,
                subtotal_amount,
                discount_amount,
                total_amount,
                status,
                _clean_text(payload.get("notes")),
            ),
        )
        quote_id = cursor.lastrowid
        for item in items:
            connection.execute(
                """
                INSERT INTO quote_items (quote_id, product_id, item_name, unit, quantity, unit_price, total_price)
                VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    quote_id,
                    item["product_id"],
                    item["item_name"],
                    item["unit"],
                    item["quantity"],
                    item["unit_price"],
                    item["total_price"],
                ),
            )
    return next(quote for quote in list_quotes() if quote["id"] == quote_id)


def update_quote(quote_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    with get_connection() as connection:
        existing = connection.execute("SELECT * FROM quotes WHERE id = ?", (quote_id,)).fetchone()
        if not existing:
            raise ServiceError("Orçamento não encontrado.", 404)

        items, subtotal_amount = _prepare_quote_items(payload.get("items") or [])
        status = payload.get("status") or existing["status"]
        if status not in QUOTE_STATUSES:
            raise ServiceError("Escolha um status de orçamento válido.")
        customer_name_manual = _clean_text(
            payload.get("customer_name_manual")
            if payload.get("customer_name_manual") is not None
            else existing.get("customer_name_manual")
        )
        discount_amount = _parse_amount(
            payload.get("discount_amount", existing["discount_amount"]),
            "discount_amount",
            min_value=0,
            allow_zero=True,
        )
        if discount_amount > subtotal_amount:
            raise ServiceError("O desconto não pode ser maior que o subtotal do orçamento.")
        total_amount = round_money(subtotal_amount - discount_amount)

        connection.execute(
            """
            UPDATE quotes
            SET quote_date = ?, validity_date = ?, customer_id = ?, customer_name_manual = ?,
                subtotal_amount = ?, discount_amount = ?, total_amount = ?, status = ?, notes = ?
            WHERE id = ?
            """,
            (
                ensure_date(payload.get("quote_date"), existing["quote_date"]),
                ensure_date(payload.get("validity_date"), existing["validity_date"] or existing["quote_date"]),
                int(payload["customer_id"]) if payload.get("customer_id") else None,
                customer_name_manual or None,
                subtotal_amount,
                discount_amount,
                total_amount,
                status,
                _clean_text(payload.get("notes") if payload.get("notes") is not None else existing["notes"]),
                quote_id,
            ),
        )
        connection.execute("DELETE FROM quote_items WHERE quote_id = ?", (quote_id,))
        for item in items:
            connection.execute(
                """
                INSERT INTO quote_items (quote_id, product_id, item_name, unit, quantity, unit_price, total_price)
                VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    quote_id,
                    item["product_id"],
                    item["item_name"],
                    item["unit"],
                    item["quantity"],
                    item["unit_price"],
                    item["total_price"],
                ),
            )
    return next(quote for quote in list_quotes() if quote["id"] == quote_id)


def delete_quote(quote_id: int) -> None:
    with get_connection() as connection:
        deleted = connection.execute("DELETE FROM quotes WHERE id = ?", (quote_id,)).rowcount
    if not deleted:
        raise ServiceError("Orçamento não encontrado.", 404)


def _prepare_sale_items(connection: Any, raw_items: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], float]:
    if not raw_items:
        raise ServiceError("Adicione pelo menos um item na venda.")

    product_ids: list[int] = []
    normalized_rows: list[dict[str, Any]] = []
    for index, raw_item in enumerate(raw_items, start=1):
        product_id = _parse_int(raw_item.get("product_id"), f"product_id_{index}", min_value=1)
        quantity = _parse_amount(raw_item.get("quantity"), f"quantity_{index}", min_value=0.01, allow_zero=False)
        unit_price = _parse_amount(raw_item.get("unit_price"), f"unit_price_{index}", min_value=0, allow_zero=True)
        product_ids.append(product_id)
        normalized_rows.append(
            {
                "product_id": product_id,
                "quantity": quantity,
                "unit_price": unit_price,
                "description": _clean_text(raw_item.get("description")),
            }
        )

    products_map = _load_products_map(connection, product_ids)
    items: list[dict[str, Any]] = []
    total_amount = 0.0
    for index, item in enumerate(normalized_rows, start=1):
        product = products_map.get(item["product_id"])
        if not product:
            raise ServiceError(f"O produto do item {index} não foi encontrado.")

        quantity = item["quantity"]
        unit_price = item["unit_price"]
        subtotal = round_money(quantity * unit_price)
        items.append(
            {
                "product_id": product["id"],
                "sku": product.get("sku") or product.get("code"),
                "description": item["description"] or product["name"],
                "unit": product.get("unit") or "UN",
                "quantity": quantity,
                "unit_price": unit_price,
                "total_price": subtotal,
                "ncm": _clean_text(product.get("ncm")),
                "cfop": _clean_text(product.get("cfop_default")),
                "origin": _clean_text(product.get("origin")),
                "csosn": _clean_text(product.get("csosn")),
                "product_name": product.get("name"),
                "product_code": product.get("code"),
            }
        )
        total_amount += subtotal
    return items, round_money(total_amount)


def _validate_sale_stock(connection: Any, items: list[dict[str, Any]]) -> None:
    products_map = _load_products_map(connection, [item["product_id"] for item in items])
    allow_negative = _allow_negative_stock(connection)
    for item in items:
        product = products_map.get(item["product_id"])
        available_stock = float(product.get("stock_quantity") or 0)
        if not allow_negative and available_stock < float(item["quantity"]):
            raise ServiceError(
                f"Estoque insuficiente para o produto '{product['name']}'. Disponível: {round_money(available_stock)}.",
                409,
            )


def _write_sale_items(connection: Any, sale_id: int, items: list[dict[str, Any]]) -> None:
    for item in items:
        connection.execute(
            """
            INSERT INTO sale_items (
                sale_id, product_id, sku, description, unit, quantity,
                unit_price, total_price, ncm, cfop, origin, csosn
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                sale_id,
                item["product_id"],
                item["sku"],
                item["description"],
                item["unit"],
                item["quantity"],
                item["unit_price"],
                item["total_price"],
                item["ncm"],
                item["cfop"],
                item["origin"],
                item["csosn"],
            ),
        )


def _apply_sale_stock(connection: Any, sale_id: int, items: list[dict[str, Any]], user_id: int | None = None) -> None:
    _validate_sale_stock(connection, items)
    products_map = _load_products_map(connection, [item["product_id"] for item in items])
    for item in items:
        product = products_map[item["product_id"]]
        balance_before = float(product["stock_quantity"])
        balance_after = round_money(balance_before - float(item["quantity"]))
        connection.execute(
            "UPDATE products SET stock_quantity = ?, updated_at = ? WHERE id = ?",
            (balance_after, iso_now(), item["product_id"]),
        )
        _insert_stock_movement(
            connection,
            product_id=item["product_id"],
            movement_type="SAIDA",
            quantity=item["quantity"],
            balance_before=balance_before,
            balance_after=balance_after,
            reason=f"Baixa automática pela venda #{sale_id}",
            document_reference=f"VENDA-{sale_id}",
            user_id=user_id,
        )
        product["stock_quantity"] = balance_after


def _reverse_sale_stock(connection: Any, sale_id: int, items: list[dict[str, Any]], user_id: int | None = None) -> None:
    products_map = _load_products_map(connection, [item["product_id"] for item in items if item.get("product_id")])
    for item in items:
        product_id = item.get("product_id")
        if not product_id or product_id not in products_map:
            continue
        product = products_map[product_id]
        balance_before = float(product["stock_quantity"])
        balance_after = round_money(balance_before + float(item["quantity"]))
        connection.execute(
            "UPDATE products SET stock_quantity = ?, updated_at = ? WHERE id = ?",
            (balance_after, iso_now(), product_id),
        )
        _insert_stock_movement(
            connection,
            product_id=product_id,
            movement_type="ENTRADA",
            quantity=item["quantity"],
            balance_before=balance_before,
            balance_after=balance_after,
            reason=f"Estorno de estoque da venda #{sale_id}",
            document_reference=f"VENDA-{sale_id}",
            user_id=user_id,
        )
        product["stock_quantity"] = balance_after


def _serialize_sale_rows(rows: list[dict[str, Any]], items_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    items_by_sale: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for row in items_rows:
        item = dict(row)
        items_by_sale[row["sale_id"]].append(
            {
                "id": row["id"],
                "product_id": row["product_id"],
                "sku": item.get("sku"),
                "description": item.get("description") or item.get("product_name") or "-",
                "product_name": item.get("description") or item.get("product_name") or "-",
                "product_code": item.get("product_code") or item.get("sku") or "-",
                "unit": item.get("unit") or "UN",
                "quantity": row["quantity"],
                "unit_price": row["unit_price"],
                "total_price": row["total_price"],
                "ncm": item.get("ncm"),
                "cfop": item.get("cfop"),
                "origin": item.get("origin"),
                "csosn": item.get("csosn"),
            }
        )

    sales: list[dict[str, Any]] = []
    for row in rows:
        sale = dict(row)
        sale["amount"] = round_money(sale.get("amount", sale.get("total_amount", 0)))
        sale["total_amount"] = round_money(sale.get("total_amount", sale["amount"]))
        sale["sale_time"] = sale.get("sale_time") or "08:00"
        sale["period"] = sale.get("period") or _infer_period(sale["sale_time"])
        sale["payment_method"] = _normalize_payment_method(sale.get("payment_method"), default="-") or "-"
        sale["customer_name"] = sale.get("customer_name") or "Consumidor final"
        sale["items"] = items_by_sale.get(row["id"], [])
        sale["item_count"] = len(sale["items"])
        sale["customer"] = {
            "id": sale.get("customer_id"),
            "name": sale.get("customer_name"),
            "document": sale.get("customer_document"),
            "address": sale.get("customer_address"),
            "phone": sale.get("customer_phone"),
        }
        sales.append(sale)
    return sales


def list_sales() -> list[dict[str, Any]]:
    with get_connection() as connection:
        rows = connection.execute(
            """
            SELECT
                sales.*,
                customers.name AS customer_name,
                customers.document AS customer_document,
                customers.address AS customer_address,
                customers.phone AS customer_phone
            FROM sales
            LEFT JOIN customers ON customers.id = sales.customer_id
            ORDER BY sales.sale_date DESC, sales.sale_time DESC, sales.id DESC
            """
        ).fetchall()
        sale_ids = [row["id"] for row in rows]
        if sale_ids:
            placeholders = ",".join("?" for _ in sale_ids)
            items_rows = connection.execute(
                f"""
                SELECT sale_items.*, products.name AS product_name, products.code AS product_code
                FROM sale_items
                LEFT JOIN products ON products.id = sale_items.product_id
                WHERE sale_items.sale_id IN ({placeholders})
                ORDER BY sale_items.id
                """,
                tuple(sale_ids),
            ).fetchall()
        else:
            items_rows = []
    return _serialize_sale_rows(rows, items_rows)


def get_sale(sale_id: int) -> dict[str, Any]:
    sale = next((item for item in list_sales() if item["id"] == sale_id), None)
    if not sale:
        raise ServiceError("Venda não encontrada.", 404)
    return sale


def create_sale(payload: dict[str, Any]) -> dict[str, Any]:
    sale_date = ensure_date(payload.get("sale_date"), _current_sale_date())
    sale_time = _normalize_sale_time(payload.get("sale_time"))
    payment_method = _normalize_payment_method(payload.get("payment_method"))
    if payment_method not in SALES_PAYMENT_METHODS:
        raise ServiceError("Escolha uma forma de pagamento válida para a venda.")

    customer_id = payload.get("customer_id") or None
    raw_items = payload.get("items") or []
    user_id = _parse_int(payload.get("_user_id"), "user_id", default=0) if payload.get("_user_id") else None

    with get_connection() as connection:
        if raw_items:
            items, total_amount = _prepare_sale_items(connection, raw_items)
        else:
            items = []
            total_amount = _parse_amount(payload.get("amount", payload.get("total_amount")), "amount", min_value=0.01, allow_zero=False)

        period = _infer_period(sale_time)
        cursor = connection.execute(
            """
            INSERT INTO sales (sale_date, sale_time, period, amount, customer_id, total_amount, payment_method, notes, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                sale_date,
                sale_time,
                period,
                total_amount,
                int(customer_id) if customer_id else None,
                total_amount,
                payment_method,
                _clean_text(payload.get("notes")),
                iso_now(),
            ),
        )
        sale_id = cursor.lastrowid
        if items:
            _write_sale_items(connection, sale_id, items)
            _apply_sale_stock(connection, sale_id, items, user_id)

    return get_sale(sale_id)


def update_sale(sale_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    with get_connection() as connection:
        existing_row = connection.execute("SELECT * FROM sales WHERE id = ?", (sale_id,)).fetchone()
        if not existing_row:
            raise ServiceError("Venda não encontrada.", 404)

        existing_items = connection.execute(
            """
            SELECT *
            FROM sale_items
            WHERE sale_id = ?
            ORDER BY id
            """,
            (sale_id,),
        ).fetchall()

        user_id = _parse_int(payload.get("_user_id"), "user_id", default=0) if payload.get("_user_id") else None
        sale_time = _normalize_sale_time(payload.get("sale_time", existing_row["sale_time"]))
        payment_method = _normalize_payment_method(payload.get("payment_method", existing_row["payment_method"]))
        if payment_method not in SALES_PAYMENT_METHODS:
            raise ServiceError("Escolha uma forma de pagamento válida para a venda.")

        update_items = "items" in payload
        raw_items = payload.get("items") if update_items else None
        if update_items and raw_items:
            items, total_amount = _prepare_sale_items(connection, raw_items)
        elif update_items and raw_items == []:
            items = []
            total_amount = _parse_amount(
                payload.get("amount", payload.get("total_amount", existing_row["total_amount"])),
                "amount",
                min_value=0.01,
                allow_zero=False,
            )
        else:
            items = None
            total_amount = _parse_amount(
                payload.get("amount", payload.get("total_amount", existing_row["total_amount"])),
                "amount",
                min_value=0.01,
                allow_zero=False,
            )

        if update_items and existing_items:
            _reverse_sale_stock(connection, sale_id, _rows_to_dicts(existing_items), user_id)
            connection.execute("DELETE FROM sale_items WHERE sale_id = ?", (sale_id,))

        connection.execute(
            """
            UPDATE sales
            SET sale_date = ?, sale_time = ?, period = ?, amount = ?, customer_id = ?, total_amount = ?, payment_method = ?, notes = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                ensure_date(payload.get("sale_date"), existing_row["sale_date"]),
                sale_time,
                _infer_period(sale_time),
                total_amount,
                int(payload["customer_id"]) if "customer_id" in payload and payload.get("customer_id") else (existing_row["customer_id"] if "customer_id" not in payload else None),
                total_amount,
                payment_method,
                _clean_text(payload.get("notes") if payload.get("notes") is not None else existing_row["notes"]),
                iso_now(),
                sale_id,
            ),
        )

        if update_items and items:
            _write_sale_items(connection, sale_id, items)
            _apply_sale_stock(connection, sale_id, items, user_id)

    return get_sale(sale_id)


def delete_sale(sale_id: int) -> None:
    with get_connection() as connection:
        sale = connection.execute("SELECT id FROM sales WHERE id = ?", (sale_id,)).fetchone()
        if not sale:
            raise ServiceError("Venda não encontrada.", 404)

        old_items = connection.execute("SELECT * FROM sale_items WHERE sale_id = ?", (sale_id,)).fetchall()
        if old_items:
            _reverse_sale_stock(connection, sale_id, _rows_to_dicts(old_items))
        connection.execute("DELETE FROM sales WHERE id = ?", (sale_id,))


def list_expenses() -> list[dict[str, Any]]:
    with get_connection() as connection:
        rows = connection.execute(
            """
            SELECT *
            FROM expenses
            ORDER BY payment_date DESC, id DESC
            """
        ).fetchall()
    return [_serialize_expense_row(dict(row)) for row in rows]


def get_expense(expense_id: int) -> dict[str, Any]:
    with get_connection() as connection:
        row = connection.execute("SELECT * FROM expenses WHERE id = ?", (expense_id,)).fetchone()
    if not row:
        raise ServiceError("Conta paga não encontrada.", 404)
    return _serialize_expense_row(dict(row))


def create_expense(payload: dict[str, Any]) -> dict[str, Any]:
    payment_date = ensure_date(payload.get("payment_date"))
    description = _require_text(payload.get("description"), "description")
    amount = _parse_amount(payload.get("amount"), "amount", min_value=0.01, allow_zero=False)
    payment_method = _normalize_payment_method(_require_text(payload.get("payment_method"), "payment_method"))

    if payment_method not in PAYMENT_METHODS:
        raise ServiceError("Escolha uma forma de pagamento válida.")

    with get_connection() as connection:
        cursor = connection.execute(
            """
            INSERT INTO expenses (payment_date, description, category, amount, payment_method, supplier, notes)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                payment_date,
                description,
                DEFAULT_EXPENSE_CATEGORY,
                amount,
                payment_method,
                "",
                "",
            ),
        )
        expense_id = cursor.lastrowid
    return get_expense(expense_id)


def update_expense(expense_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    current = get_expense(expense_id)
    payment_method = _normalize_payment_method(
        _require_text(payload.get("payment_method", current["payment_method"]), "payment_method"),
    )
    if payment_method not in PAYMENT_METHODS:
        raise ServiceError("Escolha uma forma de pagamento válida.")

    with get_connection() as connection:
        connection.execute(
            """
            UPDATE expenses
            SET payment_date = ?, description = ?, category = ?, amount = ?,
                payment_method = ?, supplier = ?, notes = ?
            WHERE id = ?
            """,
            (
                ensure_date(payload.get("payment_date"), current["payment_date"]),
                _require_text(payload.get("description", current["description"]), "description"),
                DEFAULT_EXPENSE_CATEGORY,
                _parse_amount(payload.get("amount", current["amount"]), "amount", min_value=0.01, allow_zero=False),
                payment_method,
                "",
                "",
                expense_id,
            ),
        )
    return get_expense(expense_id)


def delete_expense(expense_id: int) -> None:
    with get_connection() as connection:
        deleted = connection.execute("DELETE FROM expenses WHERE id = ?", (expense_id,)).rowcount
    if not deleted:
        raise ServiceError("Conta paga não encontrada.", 404)


def list_bills() -> list[dict[str, Any]]:
    with get_connection() as connection:
        rows = connection.execute(
            """
            SELECT *
            FROM bills
            ORDER BY is_paid ASC, due_date ASC, id DESC
            """
        ).fetchall()
    return [_serialize_bill(row) for row in rows]


def get_bill(bill_id: int) -> dict[str, Any]:
    with get_connection() as connection:
        row = connection.execute("SELECT * FROM bills WHERE id = ?", (bill_id,)).fetchone()
    if not row:
        raise ServiceError("Boleto nÃ£o encontrado.", 404)
    return _serialize_bill(row)


def create_bill(payload: dict[str, Any]) -> dict[str, Any]:
    beneficiary = _require_text(payload.get("beneficiary"), "beneficiary")
    due_date = ensure_date(payload.get("due_date"))
    amount = _parse_amount(payload.get("amount"), "amount", min_value=0.01, allow_zero=False)
    is_paid = _parse_bool(payload.get("is_paid"), False)
    now = iso_now()

    with get_connection() as connection:
        cursor = connection.execute(
            """
            INSERT INTO bills (beneficiary, due_date, amount, is_paid, notes, paid_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                beneficiary,
                due_date,
                amount,
                is_paid,
                _clean_text(payload.get("notes")),
                now if is_paid else None,
                now,
            ),
        )
        bill_id = cursor.lastrowid
        _sync_paid_bill_expense(
            connection,
            bill_id=bill_id,
            is_paid=is_paid,
            amount=amount,
            payment_date=_payment_date_from_paid_at(now if is_paid else None),
        )

    return get_bill(bill_id)


def update_bill(bill_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    current = get_bill(bill_id)
    next_amount = _parse_amount(payload.get("amount", current["amount"]), "amount", min_value=0.01, allow_zero=False)
    next_is_paid = _parse_bool(payload.get("is_paid"), current["is_paid"])
    now = iso_now()

    if current["is_paid"] and next_is_paid:
        paid_at = current.get("paid_at") or now
    elif not current["is_paid"] and next_is_paid:
        paid_at = now
    else:
        paid_at = None

    with get_connection() as connection:
        connection.execute(
            """
            UPDATE bills
            SET beneficiary = ?, due_date = ?, amount = ?, is_paid = ?, notes = ?, paid_at = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                _require_text(payload.get("beneficiary", current["beneficiary"]), "beneficiary"),
                ensure_date(payload.get("due_date"), current["due_date"]),
                next_amount,
                next_is_paid,
                _clean_text(payload.get("notes") if payload.get("notes") is not None else current.get("notes")),
                paid_at,
                now,
                bill_id,
            ),
        )
        _sync_paid_bill_expense(
            connection,
            bill_id=bill_id,
            is_paid=next_is_paid,
            amount=next_amount,
            payment_date=_payment_date_from_paid_at(paid_at),
        )

    return get_bill(bill_id)


def delete_bill(bill_id: int) -> None:
    with get_connection() as connection:
        connection.execute("DELETE FROM expenses WHERE linked_bill_id = ?", (bill_id,))
        deleted = connection.execute("DELETE FROM bills WHERE id = ?", (bill_id,)).rowcount
    if not deleted:
        raise ServiceError("Boleto nÃ£o encontrado.", 404)


def _serialize_missing_item_row(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": int(row["id"]),
        "name": row["name"],
    }


def list_missing_items() -> list[dict[str, Any]]:
    with get_connection() as connection:
        rows = connection.execute(
            """
            SELECT id, name
            FROM missing_items
            ORDER BY id DESC
            """
        ).fetchall()
    return [_serialize_missing_item_row(dict(row)) for row in rows]


def get_missing_item(item_id: int) -> dict[str, Any]:
    with get_connection() as connection:
        row = connection.execute(
            """
            SELECT id, name
            FROM missing_items
            WHERE id = ?
            """,
            (item_id,),
        ).fetchone()
    if not row:
        raise ServiceError("Item faltante nÃ£o encontrado.", 404)
    return _serialize_missing_item_row(dict(row))


def create_missing_item(payload: dict[str, Any]) -> dict[str, Any]:
    name = _require_text(payload.get("name"), "name")
    with get_connection() as connection:
        cursor = connection.execute(
            """
            INSERT INTO missing_items (name)
            VALUES (?)
            """,
            (name,),
        )
        item_id = cursor.lastrowid
    return get_missing_item(item_id)


def update_missing_item(item_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    name = _require_text(payload.get("name"), "name")
    with get_connection() as connection:
        updated = connection.execute(
            """
            UPDATE missing_items
            SET name = ?
            WHERE id = ?
            """,
            (name, item_id),
        ).rowcount
    if not updated:
        raise ServiceError("Item faltante nÃ£o encontrado.", 404)
    return get_missing_item(item_id)


def delete_missing_item(item_id: int) -> None:
    with get_connection() as connection:
        deleted = connection.execute("DELETE FROM missing_items WHERE id = ?", (item_id,)).rowcount
    if not deleted:
        raise ServiceError("Item faltante nÃ£o encontrado.", 404)


def list_checks() -> list[dict[str, Any]]:
    with get_connection() as connection:
        rows = connection.execute(
            """
            SELECT *
            FROM checks
            ORDER BY due_date ASC, id DESC
            """
        ).fetchall()
    return [_serialize_check(row) for row in rows]


def get_check(check_id: int) -> dict[str, Any]:
    with get_connection() as connection:
        row = connection.execute("SELECT * FROM checks WHERE id = ?", (check_id,)).fetchone()
    if not row:
        raise ServiceError("Cheque não encontrado.", 404)
    return _serialize_check(row)


def _ensure_check_number_available(connection: Any, check_number: str, *, exclude_id: int | None = None) -> None:
    if check_number == "S/N":
        return

    query = "SELECT id FROM checks WHERE check_number = ?"
    params: list[Any] = [check_number]
    if exclude_id is not None:
        query += " AND id <> ?"
        params.append(exclude_id)

    existing = connection.execute(query, tuple(params)).fetchone()
    if existing:
        raise ServiceError("Já existe um cheque com esse número.")


def create_check(payload: dict[str, Any]) -> dict[str, Any]:
    check_number = _normalize_check_number(payload.get("check_number"))
    beneficiary = _require_text(payload.get("beneficiary"), "beneficiary")
    amount = _parse_amount(payload.get("amount"), "amount", min_value=0.01, allow_zero=False)
    issue_date = ensure_date(payload.get("issue_date"))
    due_date = ensure_date(payload.get("due_date"))
    status = _require_text(payload.get("status"), "status")
    if status not in CHECK_STATUSES:
        raise ServiceError("Escolha um status de cheque válido.")
    if due_date < issue_date:
        raise ServiceError("A data prevista do cheque não pode ser anterior à data de emissão.")

    try:
        with get_connection() as connection:
            _ensure_check_number_available(connection, check_number)
            cursor = connection.execute(
                """
                INSERT INTO checks (check_number, beneficiary, amount, issue_date, due_date, status, notes)
                VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    check_number,
                    beneficiary,
                    amount,
                    issue_date,
                    due_date,
                    status,
                    _clean_text(payload.get("notes")),
                ),
            )
            check_id = cursor.lastrowid
    except UniqueViolation as exc:
        raise ServiceError("Já existe um cheque com esse número.") from exc
    return get_check(check_id)


def update_check(check_id: int, payload: dict[str, Any]) -> dict[str, Any]:
    current = get_check(check_id)
    next_check_number = _normalize_check_number(payload.get("check_number", current["check_number"]))
    status = _require_text(payload.get("status", current["status"]), "status")
    if status not in CHECK_STATUSES:
        raise ServiceError("Escolha um status de cheque válido.")
    issue_date = ensure_date(payload.get("issue_date"), current["issue_date"])
    due_date = ensure_date(payload.get("due_date"), current["due_date"])
    if due_date < issue_date:
        raise ServiceError("A data prevista do cheque não pode ser anterior à data de emissão.")

    try:
        with get_connection() as connection:
            _ensure_check_number_available(connection, next_check_number, exclude_id=check_id)
            connection.execute(
                """
                UPDATE checks
                SET check_number = ?, beneficiary = ?, amount = ?, issue_date = ?,
                    due_date = ?, status = ?, notes = ?
                WHERE id = ?
                """,
                (
                    next_check_number,
                    _require_text(payload.get("beneficiary", current["beneficiary"]), "beneficiary"),
                    _parse_amount(payload.get("amount", current["amount"]), "amount", min_value=0.01, allow_zero=False),
                    issue_date,
                    due_date,
                    status,
                    _clean_text(payload.get("notes") if payload.get("notes") is not None else current["notes"]),
                    check_id,
                ),
            )
    except UniqueViolation as exc:
        raise ServiceError("Já existe um cheque com esse número.") from exc
    return get_check(check_id)


def delete_check(check_id: int) -> None:
    with get_connection() as connection:
        deleted = connection.execute("DELETE FROM checks WHERE id = ?", (check_id,)).rowcount
    if not deleted:
        raise ServiceError("Cheque não encontrado.", 404)


def list_stock_movements() -> list[dict[str, Any]]:
    with get_connection() as connection:
        rows = connection.execute(
            """
            SELECT
                stock_movements.*,
                products.name AS product_name,
                products.sku AS product_sku
            FROM stock_movements
            JOIN products ON products.id = stock_movements.product_id
            ORDER BY stock_movements.created_at DESC, stock_movements.id DESC
            """
        ).fetchall()
    return [_serialize_stock_movement(row) for row in rows]


def create_stock_movement(payload: dict[str, Any]) -> dict[str, Any]:
    product_id = _parse_int(payload.get("product_id"), "product_id", min_value=1)
    movement_type = _clean_text(payload.get("movement_type")).upper()
    if movement_type not in STOCK_MOVEMENT_TYPES:
        raise ServiceError("Escolha um tipo de movimentação válido.")

    reason = _require_text(payload.get("reason"), "reason")
    document_reference = _clean_text(payload.get("document_reference"))
    user_id = _parse_int(payload.get("_user_id"), "user_id", default=0) if payload.get("_user_id") else None

    with get_connection() as connection:
        product_row = connection.execute("SELECT * FROM products WHERE id = ?", (product_id,)).fetchone()
        if not product_row:
            raise ServiceError("Produto não encontrado para movimentação de estoque.", 404)

        current_stock = float(product_row["stock_quantity"] or 0)
        allow_negative = _allow_negative_stock(connection)

        if movement_type == "AJUSTE":
            target_stock = _parse_amount(
                payload.get("target_stock", payload.get("quantity")),
                "target_stock",
                min_value=0,
                allow_zero=True,
            )
            quantity = round_money(target_stock - current_stock)
            next_stock = target_stock
        else:
            quantity = _parse_amount(payload.get("quantity"), "quantity", min_value=0.01, allow_zero=False)
            next_stock = current_stock + quantity if movement_type == "ENTRADA" else current_stock - quantity

        if not allow_negative and next_stock < 0:
            raise ServiceError("A movimentação deixaria o estoque negativo e isso não está liberado.", 409)

        connection.execute(
            "UPDATE products SET stock_quantity = ?, updated_at = ? WHERE id = ?",
            (round_money(next_stock), iso_now(), product_id),
        )
        _insert_stock_movement(
            connection,
            product_id=product_id,
            movement_type=movement_type,
            quantity=quantity,
            balance_before=current_stock,
            balance_after=next_stock,
            reason=reason,
            document_reference=document_reference,
            user_id=user_id,
        )

    return list_stock_movements()[0]


def get_stock_overview() -> dict[str, Any]:
    products = list_products()
    active_products = [product for product in products if product["active"]]
    low_stock = [product for product in active_products if product["low_stock"]]
    out_of_stock = [product for product in active_products if product["out_of_stock"]]
    return {
        "total_products": len(active_products),
        "low_stock_products": len(low_stock),
        "out_of_stock_products": len(out_of_stock),
        "estimated_sale_value": round_money(sum(float(product["stock_quantity"]) * float(product["sale_price"]) for product in active_products)),
    }


def _ensure_fiscal_settings_exists() -> dict[str, Any]:
    with get_connection() as connection:
        row = connection.execute("SELECT * FROM fiscal_settings ORDER BY id ASC LIMIT 1").fetchone()
        if row:
            return dict(row)
        cursor = connection.execute(
            """
            INSERT INTO fiscal_settings (
                company_name, trade_name, environment, provider_name, created_at, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (None, None, "homologation", "mock", iso_now(), iso_now()),
        )
        return get_fiscal_settings(cursor.lastrowid)


def get_fiscal_settings(settings_id: int | None = None) -> dict[str, Any]:
    _ensure_fiscal_settings_exists()
    with get_connection() as connection:
        if settings_id:
            row = connection.execute("SELECT * FROM fiscal_settings WHERE id = ?", (settings_id,)).fetchone()
        else:
            row = connection.execute("SELECT * FROM fiscal_settings ORDER BY id ASC LIMIT 1").fetchone()
    if not row:
        raise ServiceError("Configurações fiscais não encontradas.", 404)
    return dict(row)


def _validate_cnpj(value: str) -> str:
    digits = _clean_digits(value)
    if len(digits) != 14 or digits == digits[0] * 14:
        raise ServiceError("Informe um CNPJ valido com 14 digitos.")

    weights = ((5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2), (6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2))
    for index, sequence in enumerate(weights, start=12):
        total = sum(int(digits[position]) * sequence[position] for position in range(index))
        remainder = total % 11
        expected = 0 if remainder < 2 else 11 - remainder
        if expected != int(digits[index]):
            raise ServiceError("Informe um CNPJ valido.")
    return digits


def update_fiscal_settings(payload: dict[str, Any]) -> dict[str, Any]:
    current = get_fiscal_settings()
    environment = _clean_text(payload.get("environment", current["environment"])) or "homologation"
    provider_name = _clean_text(payload.get("provider_name", current.get("provider_name"))) or "mock"

    if environment not in FISCAL_ENVIRONMENTS:
        raise ServiceError("Escolha um ambiente fiscal válido.")
    if provider_name not in FISCAL_PROVIDER_OPTIONS:
        raise ServiceError("Escolha um provider fiscal válido.")

    with get_connection() as connection:
        connection.execute(
            """
            UPDATE fiscal_settings
            SET company_name = ?, trade_name = ?, cnpj = ?, state_registration = ?, tax_regime = ?,
                street = ?, number = ?, complement = ?, district = ?, city = ?, state = ?, zip_code = ?,
                phone = ?, email = ?, default_series = ?, next_nfe_number = ?, environment = ?, provider_name = ?,
                api_token = ?, api_url = ?, certificate_path = ?, certificate_password = ?, csc = ?,
                allow_negative_stock = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                _clean_text(payload.get("company_name") if payload.get("company_name") is not None else current.get("company_name")),
                _clean_text(payload.get("trade_name") if payload.get("trade_name") is not None else current.get("trade_name")),
                _validate_cnpj(payload.get("cnpj") if payload.get("cnpj") is not None else current.get("cnpj")),
                _clean_text(payload.get("state_registration") if payload.get("state_registration") is not None else current.get("state_registration")),
                _clean_text(payload.get("tax_regime") if payload.get("tax_regime") is not None else current.get("tax_regime")),
                _clean_text(payload.get("street") if payload.get("street") is not None else current.get("street")),
                _clean_text(payload.get("number") if payload.get("number") is not None else current.get("number")),
                _clean_text(payload.get("complement") if payload.get("complement") is not None else current.get("complement")),
                _clean_text(payload.get("district") if payload.get("district") is not None else current.get("district")),
                _clean_text(payload.get("city") if payload.get("city") is not None else current.get("city")),
                _clean_text(payload.get("state") if payload.get("state") is not None else current.get("state")),
                _clean_text(payload.get("zip_code") if payload.get("zip_code") is not None else current.get("zip_code")),
                _clean_text(payload.get("phone") if payload.get("phone") is not None else current.get("phone")),
                _clean_text(payload.get("email") if payload.get("email") is not None else current.get("email")),
                _parse_int(payload.get("default_series", current["default_series"]), "default_series", min_value=1),
                _parse_int(payload.get("next_nfe_number", current["next_nfe_number"]), "next_nfe_number", min_value=1),
                environment,
                provider_name,
                _clean_text(payload.get("api_token") if payload.get("api_token") is not None else current.get("api_token")),
                _clean_text(payload.get("api_url") if payload.get("api_url") is not None else current.get("api_url")),
                _clean_text(payload.get("certificate_path") if payload.get("certificate_path") is not None else current.get("certificate_path")),
                _clean_text(payload.get("certificate_password") if payload.get("certificate_password") is not None else current.get("certificate_password")),
                _clean_text(payload.get("csc") if payload.get("csc") is not None else current.get("csc")),
                _parse_bool(payload.get("allow_negative_stock"), current.get("allow_negative_stock", False)),
                iso_now(),
                current["id"],
            ),
        )
    return get_fiscal_settings(current["id"])


def _get_fiscal_settings_missing_fields(settings: dict[str, Any]) -> list[str]:
    return [
        label
        for field, label in [
            ("company_name", "Razão social"),
            ("trade_name", "Nome fantasia"),
            ("cnpj", "CNPJ"),
            ("state_registration", "Inscrição estadual"),
            ("tax_regime", "Regime tributário"),
            ("street", "Endereço"),
            ("city", "Cidade"),
            ("state", "UF"),
        ]
        if not _clean_text(settings.get(field))
    ]


def _prepare_manual_nfe_items(connection: Any, raw_items: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], float]:
    if not raw_items:
        raise ServiceError("Adicione pelo menos um item para emitir a NF-e.")

    product_ids: list[int] = []
    normalized_rows: list[dict[str, Any]] = []
    for index, raw_item in enumerate(raw_items, start=1):
        product_id = _parse_int(raw_item.get("product_id"), f"produto_{index}", min_value=1)
        quantity = _parse_amount(raw_item.get("quantity"), f"quantidade_{index}", min_value=0.01, allow_zero=False)
        unit_price = _parse_amount(raw_item.get("unit_price"), f"valor_unitario_{index}", min_value=0, allow_zero=True)
        product_ids.append(product_id)
        normalized_rows.append(
            {
                "product_id": product_id,
                "quantity": quantity,
                "unit_price": unit_price,
                "description": _clean_text(raw_item.get("description")),
            }
        )

    products_map = _load_products_map(connection, product_ids)
    items: list[dict[str, Any]] = []
    total_amount = 0.0

    for index, raw_item in enumerate(normalized_rows, start=1):
        product = products_map.get(raw_item["product_id"])
        if not product:
            raise ServiceError(f"O produto do item {index} não foi encontrado.")

        subtotal = round_money(raw_item["quantity"] * raw_item["unit_price"])
        item = {
            "product_id": product["id"],
            "sku": product.get("sku") or product.get("code"),
            "description": raw_item["description"] or product.get("name"),
            "unit": product.get("unit") or "UN",
            "quantity": raw_item["quantity"],
            "unit_price": raw_item["unit_price"],
            "total_price": subtotal,
            "ncm": _clean_text(product.get("ncm")),
            "cfop": _clean_text(product.get("cfop_default")),
            "origin": _clean_text(product.get("origin")),
            "csosn": _clean_text(product.get("csosn")),
        }

        missing_fields = [
            label
            for key, label in [
                ("description", "descrição"),
                ("ncm", "NCM"),
                ("cfop", "CFOP"),
                ("origin", "origem"),
                ("csosn", "CSOSN"),
                ("unit", "unidade"),
            ]
            if not _clean_text(item.get(key))
        ]
        if missing_fields:
            raise ServiceError(
                f"O produto '{item['description']}' não pode entrar na NF-e porque faltam: {', '.join(missing_fields)}.",
                409,
            )

        items.append(item)
        total_amount += subtotal

    return items, round_money(total_amount)


def _write_nfe_items(connection: Any, nfe_id: int, items: list[dict[str, Any]]) -> None:
    for item in items:
        connection.execute(
            """
            INSERT INTO nfe_items (
                nfe_id, product_id, sku, description, unit, quantity,
                unit_price, total_price, ncm, cfop, origin, csosn
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                nfe_id,
                item.get("product_id"),
                item.get("sku"),
                item.get("description"),
                item.get("unit") or "UN",
                item.get("quantity"),
                item.get("unit_price"),
                item.get("total_price"),
                item.get("ncm"),
                item.get("cfop"),
                item.get("origin"),
                item.get("csosn"),
            ),
        )


def _emit_nfe_document(
    connection: Any,
    *,
    latest_settings: dict[str, Any],
    sale_id: int | None,
    source_type: str,
    sale_snapshot: dict[str, Any],
) -> int:
    provider = get_fiscal_provider(latest_settings.get("provider_name"))
    xml_storage = XmlStorageService()
    danfe_service = DanfeService()
    number = int(latest_settings["next_nfe_number"])
    series = int(latest_settings["default_series"])

    provider_payload = {
        "number": number,
        "series": series,
        "sale_id": sale_id,
        "sale_date": sale_snapshot.get("sale_date"),
        "sale_time": sale_snapshot.get("sale_time"),
        "payment_method": sale_snapshot.get("payment_method"),
        "notes": sale_snapshot.get("notes"),
        "customer_name": sale_snapshot.get("customer_name"),
        "customer": sale_snapshot.get("customer") or {},
        "items": sale_snapshot.get("items") or [],
        "total_amount": sale_snapshot.get("total_amount"),
    }

    emission = provider.emit_nfe(dict(latest_settings), provider_payload)
    xml_path = xml_storage.save_xml(
        access_key=emission.access_key,
        authorization_date=emission.authorization_date,
        xml_content=emission.xml_content,
    )
    nfe_record = {
        "number_nfe": emission.number,
        "series_nfe": emission.series,
        "access_key": emission.access_key,
        "authorization_protocol": emission.authorization_protocol,
        "status_nfe": emission.status,
        "authorization_date": emission.authorization_date,
        "provider_name": emission.provider_name,
        "environment": latest_settings["environment"],
    }
    pdf_path = danfe_service.save_pdf(
        authorization_date=emission.authorization_date,
        access_key=emission.access_key,
        nfe_record=nfe_record,
        sale=sale_snapshot,
        settings=dict(latest_settings),
    )

    cursor = connection.execute(
        """
        INSERT INTO nfe_issued (
            sale_id, source_type, customer_name, customer_document, customer_address, customer_phone,
            customer_notes, payment_method, total_amount, number_nfe, series_nfe, access_key,
            authorization_protocol, status_nfe, xml_path, pdf_path, authorization_date,
            sefaz_message, provider_name, environment, created_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            sale_id,
            source_type,
            _clean_text(sale_snapshot.get("customer_name")),
            _clean_text((sale_snapshot.get("customer") or {}).get("document")),
            _clean_text((sale_snapshot.get("customer") or {}).get("address")),
            _clean_text((sale_snapshot.get("customer") or {}).get("phone")),
            _clean_text(sale_snapshot.get("notes")),
            _clean_text(sale_snapshot.get("payment_method")),
            round_money(sale_snapshot.get("total_amount") or 0),
            emission.number,
            emission.series,
            emission.access_key,
            emission.authorization_protocol,
            emission.status,
            xml_path,
            pdf_path,
            emission.authorization_date,
            emission.provider_message,
            emission.provider_name,
            latest_settings["environment"],
            iso_now(),
        ),
    )
    nfe_id = cursor.lastrowid
    _write_nfe_items(connection, nfe_id, sale_snapshot.get("items") or [])

    connection.execute(
        """
        UPDATE fiscal_settings
        SET next_nfe_number = ?, updated_at = ?
        WHERE id = ?
        """,
        (number + 1, iso_now(), latest_settings["id"]),
    )
    return nfe_id


def list_nfe_issued() -> list[dict[str, Any]]:
    with get_connection() as connection:
        rows = connection.execute(
            """
            SELECT
                nfe_issued.*,
                sales.total_amount AS sale_total_amount,
                COALESCE(nfe_issued.customer_name, customers.name) AS customer_name_resolved,
                CASE
                    WHEN nfe_issued.sale_id IS NOT NULL AND COALESCE(nfe_issued.total_amount, 0) = 0
                        THEN COALESCE(sales.total_amount, 0)
                    ELSE COALESCE(nfe_issued.total_amount, sales.total_amount, 0)
                END AS total_amount_resolved,
                COALESCE(nfe_issued.payment_method, sales.payment_method, '') AS payment_method_resolved
            FROM nfe_issued
            LEFT JOIN sales ON sales.id = nfe_issued.sale_id
            LEFT JOIN customers ON customers.id = sales.customer_id
            ORDER BY COALESCE(nfe_issued.authorization_date, nfe_issued.created_at) DESC, nfe_issued.id DESC
            """
        ).fetchall()
    return [_serialize_nfe_row(row) for row in rows]


def get_nfe_issued(nfe_id: int) -> dict[str, Any]:
    issue = next((item for item in list_nfe_issued() if item["id"] == nfe_id), None)
    if not issue:
        raise ServiceError("NF-e não encontrada.", 404)
    return issue


def validate_sale_for_nfe(sale_id: int) -> dict[str, Any]:
    sale = get_sale(sale_id)
    settings = get_fiscal_settings()
    issues: list[str] = []

    company_missing = _get_fiscal_settings_missing_fields(settings)
    if company_missing:
        issues.append(f"Configurações fiscais incompletas: {', '.join(company_missing)}.")

    if not sale.get("customer_id"):
        issues.append("Selecione um cliente na venda para emitir NF-e.")

    customer = sale.get("customer") or {}
    if sale.get("customer_id"):
        try:
            customer = get_customer(int(sale["customer_id"]))
        except ServiceError as exc:
            issues.append(exc.message)
            customer = sale.get("customer") or {}
        else:
            missing_customer_fields = _get_customer_missing_fiscal_fields(customer)
            if missing_customer_fields:
                issues.append(
                    "O cliente selecionado precisa completar os dados fiscais: "
                    + ", ".join(missing_customer_fields)
                    + "."
                )

    if not sale.get("items"):
        issues.append("A venda precisa ter itens para emissão da NF-e.")

    item_checks: list[dict[str, Any]] = []
    products = {product["id"]: product for product in list_products()}
    for item in sale.get("items", []):
        missing_fields = [
            label
            for key, label in [
                ("description", "descrição"),
                ("ncm", "NCM"),
                ("cfop", "CFOP"),
                ("origin", "origem"),
                ("csosn", "CSOSN"),
                ("unit", "unidade"),
            ]
            if not _clean_text(item.get(key))
        ]
        product = products.get(item.get("product_id"))
        stock_ok = product is None or float(product.get("stock_quantity") or 0) >= 0
        if not stock_ok:
            missing_fields.append("estoque")
        item_checks.append(
            {
                "product_id": item.get("product_id"),
                "sku": item.get("sku"),
                "description": item.get("description"),
                "ok": not missing_fields,
                "missing_fields": missing_fields,
                "stock_ok": stock_ok,
            }
        )

    if any(not item["ok"] for item in item_checks):
        issues.append("Existem itens da venda com dados fiscais incompletos.")

    existing_authorized = next(
        (record for record in list_nfe_issued() if record["sale_id"] == sale_id and record["status_nfe"] == "AUTORIZADA"),
        None,
    )
    if existing_authorized:
        issues.append(f"A venda já possui NF-e autorizada (número {existing_authorized['number_nfe']}).")

    summary_checks = [
        {"label": "Configurações fiscais da empresa", "ok": not company_missing, "message": "Dados do emitente prontos." if not company_missing else ", ".join(company_missing)},
        {"label": "Cliente da venda", "ok": bool(sale.get("customer_id")), "message": sale.get("customer_name") or "Cliente não informado"},
        {"label": "Documento do cliente", "ok": bool(_clean_text(customer.get("document"))), "message": customer.get("document") or "Cliente sem CPF/CNPJ"},
        {"label": "Itens da venda", "ok": bool(sale.get("items")), "message": f"{len(sale.get('items', []))} item(ns)"},
        {"label": "Campos fiscais dos itens", "ok": not any(not item["ok"] for item in item_checks), "message": "Todos os itens possuem NCM/CFOP/origem/CSOSN."},
        {"label": "Duplicidade de NF-e", "ok": existing_authorized is None, "message": "Sem NF-e autorizada anterior." if existing_authorized is None else "Venda já autorizada anteriormente."},
    ]

    return {
        "sale_id": sale_id,
        "can_emit": not issues,
        "issues": issues,
        "checks": summary_checks,
        "items": item_checks,
        "sale": sale,
        "settings": {
            "environment": settings.get("environment"),
            "provider_name": settings.get("provider_name"),
        },
    }


def emit_sale_nfe(sale_id: int) -> dict[str, Any]:
    validation = validate_sale_for_nfe(sale_id)
    if not validation["can_emit"]:
        raise ServiceError("Não é possível emitir a NF-e: " + " ".join(validation["issues"]), 409)

    sale = validation["sale"]
    sale_snapshot = {
        "sale_date": sale["sale_date"],
        "sale_time": sale["sale_time"],
        "payment_method": sale["payment_method"],
        "notes": sale.get("notes"),
        "customer_name": sale["customer_name"],
        "customer": sale.get("customer") or {},
        "items": sale.get("items") or [],
        "total_amount": sale["total_amount"],
    }

    with get_connection() as connection:
        latest_settings = connection.execute("SELECT * FROM fiscal_settings ORDER BY id ASC LIMIT 1").fetchone()
        if not latest_settings:
            raise ServiceError("Configurações fiscais não encontradas.", 404)
        if sale.get("customer_id"):
            sale_snapshot["customer"] = _load_customer_for_nfe(connection, int(sale["customer_id"]))
        nfe_id = _emit_nfe_document(
            connection,
            latest_settings=dict(latest_settings),
            sale_id=sale_id,
            source_type="sale",
            sale_snapshot=sale_snapshot,
        )

    return get_nfe_issued(nfe_id)


def emit_manual_nfe(payload: dict[str, Any]) -> dict[str, Any]:
    customer_id = _parse_int(payload.get("customer_id"), "customer_id", min_value=1)
    payment_method = _normalize_payment_method(payload.get("payment_method")) or "Dinheiro"
    if payment_method not in PAYMENT_METHODS:
        raise ServiceError("Escolha uma forma de pagamento válida para a NF-e.")

    with get_connection() as connection:
        latest_settings = connection.execute("SELECT * FROM fiscal_settings ORDER BY id ASC LIMIT 1").fetchone()
        if not latest_settings:
            raise ServiceError("Configurações fiscais não encontradas.", 404)

        company_missing = _get_fiscal_settings_missing_fields(dict(latest_settings))
        if company_missing:
            raise ServiceError(
                "Complete as configurações fiscais antes de emitir a NF-e: " + ", ".join(company_missing) + ".",
                409,
            )

        items, total_amount = _prepare_manual_nfe_items(connection, payload.get("items") or [])
        customer = _load_customer_for_nfe(connection, customer_id)
        sale_snapshot = {
            "sale_date": ensure_date(payload.get("sale_date"), today_iso()),
            "sale_time": _normalize_sale_time(payload.get("sale_time")),
            "payment_method": payment_method,
            "notes": _clean_text(payload.get("notes")),
            "customer_name": customer["name"],
            "customer": customer,
            "items": items,
            "total_amount": total_amount,
        }
        nfe_id = _emit_nfe_document(
            connection,
            latest_settings=dict(latest_settings),
            sale_id=None,
            source_type="manual",
            sale_snapshot=sale_snapshot,
        )

    return get_nfe_issued(nfe_id)


def emit_nfe(payload: dict[str, Any]) -> dict[str, Any]:
    if payload.get("sale_id") not in (None, ""):
        sale_id = _parse_int(payload.get("sale_id"), "sale_id", min_value=1)
        return emit_sale_nfe(sale_id)
    return emit_manual_nfe(payload)


def get_nfe_file_info(nfe_id: int, file_type: str) -> dict[str, Any]:
    if file_type not in {"xml", "pdf"}:
        raise ServiceError("Tipo de arquivo fiscal inválido.", 400)
    with get_connection() as connection:
        row = connection.execute("SELECT * FROM nfe_issued WHERE id = ?", (nfe_id,)).fetchone()
    if not row:
        raise ServiceError("NF-e não encontrada.", 404)
    path_key = "xml_path" if file_type == "xml" else "pdf_path"
    file_path = _clean_text(row.get(path_key))
    if not file_path:
        raise ServiceError(f"Arquivo {file_type.upper()} ainda não foi gerado.", 404)
    return {
        "path": file_path,
        "filename": f"nfe-{row['number_nfe']}-{row['access_key']}.{file_type}",
    }


def get_today_due_checks_summary() -> dict[str, Any]:
    today = today_iso()
    with get_connection() as connection:
        due_today_rows = connection.execute(
            """
            SELECT amount
            FROM checks
            WHERE due_date = ?
              AND status NOT IN ('Compensado', 'Cancelado')
            """,
            (today,),
        ).fetchall()
        overdue_rows = connection.execute(
            """
            SELECT amount
            FROM checks
            WHERE due_date <> ?
              AND (due_date < ? OR status = 'Atrasado')
              AND status NOT IN ('Compensado', 'Cancelado')
            """,
            (today, today),
        ).fetchall()
    total_amount = round_money(sum(float(row["amount"] or 0) for row in due_today_rows))
    overdue_total_amount = round_money(sum(float(row["amount"] or 0) for row in overdue_rows))
    count = len(due_today_rows)
    overdue_count = len(overdue_rows)
    return {
        "id": f"daily-check-alert-{today}-{count}-{overdue_count}-{total_amount}-{overdue_total_amount}",
        "date": today,
        "count": count,
        "total_amount": total_amount,
        "overdue_count": overdue_count,
        "overdue_total_amount": overdue_total_amount,
        "has_alert": total_amount > 0 or overdue_total_amount > 0,
    }


def get_today_due_bills_summary() -> dict[str, Any]:
    today = today_iso()
    with get_connection() as connection:
        due_today_rows = connection.execute(
            """
            SELECT amount
            FROM bills
            WHERE due_date = ?
              AND is_paid = FALSE
            """,
            (today,),
        ).fetchall()
        overdue_rows = connection.execute(
            """
            SELECT amount
            FROM bills
            WHERE due_date < ?
              AND is_paid = FALSE
            """,
            (today,),
        ).fetchall()
    total_amount = round_money(sum(float(row["amount"] or 0) for row in due_today_rows))
    overdue_total_amount = round_money(sum(float(row["amount"] or 0) for row in overdue_rows))
    count = len(due_today_rows)
    overdue_count = len(overdue_rows)
    return {
        "id": f"daily-bill-alert-{today}-{count}-{overdue_count}-{total_amount}-{overdue_total_amount}",
        "date": today,
        "count": count,
        "total_amount": total_amount,
        "overdue_count": overdue_count,
        "overdue_total_amount": overdue_total_amount,
        "has_alert": total_amount > 0 or overdue_total_amount > 0,
    }


def get_bootstrap_data() -> dict[str, Any]:
    return {
        "products": list_products(),
        "customers": list_customers(),
        "sales": list_sales(),
        "quotes": list_quotes(),
        "expenses": list_expenses(),
        "bills": list_bills(),
        "missing_items": list_missing_items(),
        "checks": list_checks(),
        "stock_overview": get_stock_overview(),
        "nfe_issued": list_nfe_issued(),
        "fiscal_settings": get_fiscal_settings(),
        "daily_bill_alert": get_today_due_bills_summary(),
        "daily_check_alert": get_today_due_checks_summary(),
        "options": {
            "payment_methods": PAYMENT_METHODS,
            "sales_payment_methods": SALES_PAYMENT_METHODS,
            "product_units": PRODUCT_UNITS,
            "quote_item_units": QUOTE_ITEM_UNITS,
            "quote_statuses": QUOTE_STATUSES,
            "bill_statuses": BILL_STATUSES,
            "check_statuses": CHECK_STATUSES,
            "customer_person_types": CUSTOMER_PERSON_TYPES,
            "customer_ie_indicators": CUSTOMER_IE_INDICATORS,
            "stock_movement_types": STOCK_MOVEMENT_TYPES,
            "fiscal_environments": FISCAL_ENVIRONMENTS,
            "fiscal_provider_options": FISCAL_PROVIDER_OPTIONS,
        },
    }


def create_user(data: dict[str, Any]) -> dict[str, Any]:
    username = _require_text(data.get("username"), "username")
    full_name = _clean_text(data.get("full_name"))
    password = _require_text(data.get("password"), "password")
    password_hash = hash_password(password)

    try:
        with get_connection() as connection:
            cursor = connection.execute(
                """
                INSERT INTO users (username, full_name, password_hash)
                VALUES (?, ?, ?)
                """,
                (username, full_name, password_hash),
            )
            user_id = cursor.lastrowid
    except UniqueViolation as exc:
        raise ServiceError("Usuário já existe.") from exc

    return get_user_by_id(user_id)
